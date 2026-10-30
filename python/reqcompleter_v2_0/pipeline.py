"""Business pipeline, preserving v2.0 prompts with explicit task boundaries."""
from __future__ import annotations

from collections import defaultdict
import io
import json
import re
import zipfile

from . import ADAPTER_VERSION, BUSINESS_VERSION, SOURCE_SHA256
from .crud import (build_crud_matrix_df, combine_gap_detection, name_list,
                   parse_crud_triplets, parse_llm_gap_response,
                   programmatic_crud_gap_detection, clean_entity_name,
                   is_valid_entity_name)
from .protocol import TaskError


def state_json(state):
    return {
        **state,
        "crud_triples": [{"entity": entity, "useCase": use_case, "operation": operation}
                         for entity, use_case, operation in state["crud_triples"]],
    }


def save_document(store, result):
    """Requirements content only; debugging and verification material stays separate."""
    from docx import Document
    document = Document()
    document.add_heading("软件需求说明书", 0)
    if not result["converged"]:
        document.add_paragraph("本需求说明书为待复核稿：功能逻辑检查尚未判定收敛。")
    sections = [
        ("1. 软件概述", result["software_intro"]),
        ("2. 数据实体", "、".join(result["data_entities"])),
        ("3. 业务用例", "、".join(result["use_cases"])),
        ("4. 用例说明", result["use_case_description"]),
        ("5. 功能需求", result["functional_requirements"]),
        ("6. 外部接口需求", result["external_interface"]),
        ("7. 非功能需求", result["nonfunctional_requirements"]),
    ]
    for heading, text in sections:
        document.add_heading(heading, level=1)
        for paragraph in str(text).splitlines():
            if paragraph.strip():
                document.add_paragraph(paragraph)
    content = io.BytesIO()
    document.save(content)
    # Validate the generated package before publishing an atomically written artifact.
    content.seek(0)
    with zipfile.ZipFile(content) as package:
        if package.testzip() or "word/document.xml" not in package.namelist():
            raise TaskError("DOCX_FAILED", "Generated document is not a valid DOCX package")
    store.write_bytes("final/requirements.docx", content.getvalue())
    return "final/requirements.docx"


class Pipeline:
    def __init__(self, client, store, max_iterations=5):
        # Imports are lazy so --probe and protocol smoke do not run business modules.
        from .ExtraTest import DataEntityAgentWithMemory
        from .TestCurdNew import ECNUAICurdAgentWithMemory
        from .UC_Agent import ECNUAIUCAgentWithMemory
        from .ERcheckAgent import ECNUAIModelAgentWithMemory
        from .SRS_Agent import ECNUAISRSAgentWithMemory
        self.client, self.store = client, store
        self.max_iterations = max_iterations
        self.entities = DataEntityAgentWithMemory(client)
        self.crud = ECNUAICurdAgentWithMemory(client)
        self.uc = ECNUAIUCAgentWithMemory(client)
        self.er = ECNUAIModelAgentWithMemory(client)
        self.srs = ECNUAISRSAgentWithMemory(client)
        self.iterations = 0
        self.check_count = 0

    def step(self, name, function, relative, iteration=0):
        self.client.set_step(name, iteration)
        value = function()
        self.store.write_json(relative, {"step": name, "iteration": iteration, "output": value})
        self.client.emit("step", step=name, iteration=iteration, status="completed", path=relative)
        return value

    def snapshot(self, directory, name, state):
        self.store.write_json(f"{directory}/{name}.json", state_json(state))
        self.store.matrix(f"{directory}/{name}-crud-matrix", build_crud_matrix_df(state["crud_triples"]))

    def check(self, state, directory, iteration):
        self.check_count += 1
        self.client.set_step("crud_check", iteration)
        complete, programmatic = programmatic_crud_gap_detection(state["crud_triples"], state["data_entities"])
        self.store.write_json(f"{directory}/check-programmatic.json", {
            "complete": complete, "missing": programmatic,
        })
        raw = self.crud.check_uc_ett(state["data_entities"], state["use_cases"], state["crud_triples"])
        self.store.write_text(f"{directory}/check-llm.raw.txt", raw)
        llm = parse_llm_gap_response(raw)
        if not llm:
            llm = {"status": "incomplete" if programmatic else "complete", "missing": []}
        self.store.write_json(f"{directory}/check-llm.json", llm)
        combined = combine_gap_detection(llm, programmatic)
        converged = len(combined) == 0 and llm.get("status") == "complete"
        if not converged and not combined:
            combined = programmatic
        gap_map = {}
        for gap in combined:
            key = (gap["entity"].strip().lower(), gap["operation"].strip().upper())
            if key not in gap_map:
                gap_map[key] = gap
        gaps = list(gap_map.values())
        self.store.write_json(f"{directory}/gaps.json", gaps)
        self.client.emit("step", step="crud_check", iteration=iteration, status="completed",
                         converged=converged, gapCount=len(gaps))
        return converged, gaps

    def run(self, software_intro):
        state = {"software_intro": software_intro, "data_entities": [], "use_cases": [],
                 "use_case_description": "", "er_model": "", "crud_triples": []}
        raw = self.step("extract_entities", lambda: self.entities.extract_data_entities(software_intro), "initial/01-entities.raw.json")
        state["data_entities"] = name_list(raw)
        self.store.write_json("initial/entities.json", state["data_entities"])
        raw = self.step("generate_use_cases", lambda: self.entities.generate_use_cases(software_intro, state["data_entities"]), "initial/02-use-cases.raw.json")
        state["use_cases"] = name_list(raw)
        self.store.write_json("initial/use-cases.json", state["use_cases"])
        self.uc.set_data_entities(state["data_entities"])
        self.uc.set_use_cases(state["use_cases"])
        state["use_case_description"] = self.step("describe_use_cases", lambda: self.uc.generate_use_case(software_intro, state["data_entities"], state["use_cases"]), "initial/03-use-case-description.json")
        state["er_model"] = self.step("generate_er_model", lambda: self.er.generate_er_model(state["use_cases"]), "initial/04-er-model.json")
        raw = self.step("generate_crud", lambda: self.crud.generate_curd_triples(state["data_entities"], state["use_cases"], state["use_case_description"]), "initial/05-crud.raw.json")
        state["crud_triples"] = parse_crud_triplets(raw)
        state["er_model"] = self.step("check_er_model", lambda: self.er.check_generate_er_model(software_intro, state["data_entities"], state["use_cases"], state["er_model"]), "initial/06-er-checked.json")
        self.er.set_er_model(state["er_model"])
        self.snapshot("initial", "state", state)
        initial_uc_description = state["use_case_description"]

        converged = False
        for iteration in range(1, self.max_iterations + 1):
            self.iterations = iteration
            directory = f"iterations/{iteration:03d}"
            self.snapshot(directory, "before", state)
            converged, gaps = self.check(state, directory, iteration)
            if converged:
                self.snapshot(directory, "after", state)
                break
            operations = defaultdict(set)
            for gap in gaps:
                operations[gap["entity"]].add(gap["operation"])
            names = []
            verbs = {"C": "创建", "R": "查看", "U": "更新", "D": "删除"}
            for entity, missing in operations.items():
                if missing in ({"C", "R", "U", "D"}, {"R", "U", "D"}, {"U", "D"}):
                    name = f"管理{entity}"
                else:
                    name = f"{'、'.join(verbs.get(op, '') for op in sorted(missing) if op in verbs)}{entity}"
                names.append(name)
            hints = {}
            for gap in gaps:
                hints[gap["entity"]] = gap["operation"]
            self.store.write_json(f"{directory}/operation-hints.json", hints)
            if names:
                state["use_case_description"] = self.step("complete_use_cases", lambda: self.uc.generate_new_use_case(software_intro, state["data_entities"], names, operation_hint=hints), f"{directory}/use-case-description.json", iteration)
                self.uc.simple_uc = state["use_case_description"]
                for name in names:
                    if name not in state["use_cases"]:
                        state["use_cases"].append(name)
            new_uc_content = self.uc.get_new_uc() if hasattr(self.uc, 'get_new_uc') else ""
            if new_uc_content or state["use_case_description"] != initial_uc_description:
                state["er_model"] = self.step("complete_er_model", lambda: self.er.Complete_ER_Diagram(state["er_model"], state["use_case_description"]), f"{directory}/er-model.json", iteration)
                self.er.set_er_model(state["er_model"])
                match = re.search(r"(?:New Entities|新增实体)[:\s]*([^\n]+)", state["er_model"], re.IGNORECASE)
                if match:
                    raw_entities = [item.strip() for item in re.split(r"[,，、]", match.group(1)) if item.strip()]
                    new_entities = [clean_entity_name(item) for item in raw_entities]
                    for entity in new_entities:
                        if is_valid_entity_name(entity):
                            existing = {item.strip().lower() for item in state["data_entities"]}
                            if entity.lower() not in existing:
                                state["data_entities"].append(entity)
            report = json.dumps({"status": "incomplete", "missing": gaps}, ensure_ascii=False)
            raw = self.step("complete_crud", lambda: self.crud.Complete_CRUD_Triplets(state["data_entities"], state["use_case_description"], state["crud_triples"], report), f"{directory}/crud.raw.json", iteration)
            triples = parse_crud_triplets(raw)
            seen = {(entity.strip().lower(), case.strip().lower(), op.strip().upper()) for entity, case, op in state["crud_triples"]}
            for triple in triples:
                key = (triple[0].strip().lower(), triple[1].strip().lower(), triple[2].strip().upper())
                if key not in seen:
                    state["crud_triples"].append(triple)
                    seen.add(key)
            self.snapshot(directory, "after", state)
        functional = self.step("functional_requirements", lambda: self.uc.generate_functional_requirements(software_intro, state["er_model"], state["use_case_description"]), "intermediate/functional-requirements.json", self.iterations)
        external = self.step("external_interface", lambda: self.srs.generate_external_interface(software_intro, functional), "intermediate/external-interface.json", self.iterations)
        nonfunctional = self.step("nonfunctional_requirements", lambda: self.srs.generate_nonfunctional_requirement(software_intro, state["er_model"], state["use_case_description"]), "intermediate/nonfunctional-requirements.json", self.iterations)
        uc_diagram = self.step("use_case_plantuml", lambda: self.uc.generate_uc_diagram(state["use_case_description"]), "intermediate/use-case-diagram.json", self.iterations)
        er_diagram = self.step("er_plantuml", lambda: self.er.generate_er_code(state["er_model"]), "intermediate/er-diagram.json", self.iterations)
        self.store.write_text("intermediate/diagrams/use-case.puml", uc_diagram)
        self.store.write_text("intermediate/diagrams/er.puml", er_diagram)
        self.snapshot("final", "state", state)
        result = {**state_json(state), "business_version": BUSINESS_VERSION,
                  "adapter_version": ADAPTER_VERSION, "source_sha256": SOURCE_SHA256,
                  "functional_requirements": functional, "external_interface": external,
                  "nonfunctional_requirements": nonfunctional, "iterations": self.iterations,
                  "check_count": self.check_count, "converged": converged,
                  "reached_iteration_limit": self.iterations >= self.max_iterations,
                  "execution_status": "succeeded", "document_status": "pending",
                  "convergence_status": "converged" if converged else "iteration_limit"}
        self.client.set_step("save_docx", self.iterations)
        result["docx_path"] = save_document(self.store, result)
        result["document_status"] = "succeeded"
        result["files"] = sorted(self.store.files | {"final/result.json"})
        self.store.write_json("final/result.json", result)
        self.client.emit("artifact", path=result["docx_path"], kind="requirements_docx")
        return result
